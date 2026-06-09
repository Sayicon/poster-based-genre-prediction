"""
Çalıştır: python build_vize_report.py
Çıktı   : VizRaporu_231307121_MustafaKeremCekici.docx
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUT = "VizRaporu_231307121_MustafaKeremCekici.docx"

# ── yardımcılar ─────────────────────────────────────────────────────────────

def set_col_width(table, col_idx, width_cm):
    for row in table.rows:
        row.cells[col_idx].width = Cm(width_cm)

def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p

def add_para(doc, text, bold=False, italic=False, size=11, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(space_after)
    return p

def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for run in hdr[i].paragraphs[0].runs:
            run.bold = True
        hdr[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r_idx, row_data in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row_data):
            cells[c_idx].text = str(val)
            cells[c_idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if col_widths:
        for i, w in enumerate(col_widths):
            set_col_width(table, i, w)
    doc.add_paragraph()

def try_add_figure(doc, path, caption, width=5.5):
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(width))
        last = doc.paragraphs[-1]
        last.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = doc.add_paragraph(caption)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9)
    else:
        p = doc.add_paragraph(f"[Şekil bulunamadı: {path}]")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

# ── belge ────────────────────────────────────────────────────────────────────

doc = Document()

# Kenar boşlukları
for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3.0)
    section.right_margin  = Cm(2.5)

# ── kapak ───────────────────────────────────────────────────────────────────

doc.add_paragraph()
t = doc.add_paragraph("KOCAELİ ÜNİVERSİTESİ")
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
t.runs[0].bold = True
t.runs[0].font.size = Pt(14)

t2 = doc.add_paragraph("Bilişim Sistemleri Mühendisliği Bölümü")
t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
t2.runs[0].font.size = Pt(12)

doc.add_paragraph()

title = doc.add_paragraph("Veri Toplama ve Temizleme Raporu")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.runs[0].bold = True
title.runs[0].font.size = Pt(16)

subtitle = doc.add_paragraph("Film Afişlerinden Çoklu-Etiketli Tür Tahmini Projesi")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].font.size = Pt(12)

doc.add_paragraph()

info_lines = [
    "Öğrenci: Mustafa Kerem Çekici",
    "Öğrenci No: 231307121",
    "E-posta: keremcek70@gmail.com",
    "Tarih: Haziran 2026",
]
for line in info_lines:
    p = doc.add_paragraph(line)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(11)

doc.add_page_break()

# ── 1. Giriş ────────────────────────────────────────────────────────────────

add_heading(doc, "1. Giriş", 1)
add_para(doc,
    "Bu rapor, film afişlerinden çoklu-etiketli tür tahmini projesinin veri toplama ve "
    "temizleme aşamalarını kapsamlı biçimde açıklamaktadır. Proje, yalnızca poster görselini "
    "girdi alarak bir filmin türlerini (Action, Drama, Horror, vb.) tahmin eden bir model "
    "geliştirmeyi amaçlamaktadır. Bir film birden çok türe ait olabildiğinden problem doğal "
    "olarak çoklu-etiket (multi-label) sınıflandırmasıdır."
)
add_para(doc,
    "Veri toplama süreci iki temel aşamadan oluşmuştur: (1) HTML tabanlı web kazıma girişimi "
    "ve yaşanan ban sorunu, (2) TMDB resmî API'sine geçiş ve dengeli veri kümesi oluşturma. "
    "Bu rapor her iki aşamayı, karşılaşılan sorunları, alınan kararları ve nihai veri setinin "
    "özelliklerini ayrıntılı olarak sunmaktadır."
)

# ── 2. HTML Kazıma Girişimi ──────────────────────────────────────────────────

add_heading(doc, "2. HTML Kazıma Girişimi ve Yaşanan Sorunlar", 1)

add_heading(doc, "2.1 Yöntem", 2)
add_para(doc,
    "İlk veri toplama yaklaşımında themoviedb.org web sitesinin HTML sayfaları Scrapy "
    "çerçevesiyle kazınmıştır. Bot tespitini aşmak için curl_cffi kütüphanesi aracılığıyla "
    "Chrome 124 TLS parmak izi taklidi yapılmış, VPN kullanılmış ve üssel geri çekilme "
    "(exponential backoff) stratejisi uygulanmıştır. Bu yaklaşımla yaklaşık 16.300 film "
    "ve poster toplanmıştır (v1 veri kümesi, labels.csv)."
)

add_heading(doc, "2.2 Ban Sorunu", 2)
add_para(doc,
    "Site agresif hız sınırlaması ve bot koruması uyguladığından, tüm önlemlere rağmen "
    "sürekli HTTP 429 (Too Many Requests) yanıtları alınmıştır. Pratik çekim hızı "
    "3 saatte ~400–500 film düzeyinde kalıp ardından ban gelmiştir. Aşağıdaki tablo "
    "yaşanan aksaklıkları özetlemektedir:"
)

add_table(doc,
    ["Kazıma Oturumu", "HTTP 429 Sayısı", "Kalıcı Kayıp"],
    [
        ["targeted (hedefli) oturumu", "9.619", "864 film"],
        ["tmdb listing oturumu", "791", "—"],
        ["tmdb_id oturumu", "381", "+ bağlantı sıfırlama hataları"],
    ],
    col_widths=[7, 5, 5]
)

add_heading(doc, "2.3 Scraper Kod Sorunları", 2)
add_para(doc,
    "Ban sorununu ağırlaştıran teknik hatalar da tespit edilmiştir:"
)
items = [
    "CurlCffiMiddleware isteği senkron/bloklayıcı atıyordu → Twisted reaktörü donuyor, "
    "gerçek eşzamanlılık yoktu, çekim çok yavaşladı.",
    "Backoff için time.sleep() kullanımı reaktörü bloke ediyordu.",
    "targeted_spider özel ayarları ana spider'dan daha agresifti (gecikme 3 sn vs 8 sn, "
    "eşzamanlı istek 2 vs 1) → daha hızlı ban.",
    "targeted_spider, sıralı film ID'lerini gezip her sayfayı indirdikten sonra tür filtresi "
    "uyguluyordu; isteklerin ~%90'ı gereksizdi. Nadir türleri (History vb.) doldurmak "
    "bu şekilde imkânsızdı.",
]
for item in items:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(item).font.size = Pt(11)
doc.add_paragraph()

# ── 3. TMDB API'ye Geçiş ────────────────────────────────────────────────────

add_heading(doc, "3. TMDB Resmî API'ye Geçiş", 1)

add_heading(doc, "3.1 API Özellikleri", 2)
add_para(doc,
    "HTML kazıma sorunlarının kökünde web sitesinin bot koruması yattığından, veri "
    "toplama TMDB'nin resmî JSON API'sine taşınmıştır. Akademik danışman (hoca) bu "
    "yaklaşımı onaylamıştır. API'nin temel özellikleri:"
)
add_table(doc,
    ["Özellik", "Değer"],
    [
        ["Ücret", "Ücretsiz (kâr amacı gütmeyen / akademik kullanım)"],
        ["Günlük limit", "Yok (Aralık 2019'dan itibaren kaldırıldı)"],
        ["Hız", "~40–50 istek/saniye, IP başına ~20 bağlantı"],
        ["Veri formatı", "JSON; genre_ids + poster_path tek istekte"],
        ["Endpoint", "/discover/movie?with_genres=<id>&vote_count.gte=N"],
    ],
    col_widths=[6, 10]
)

add_heading(doc, "3.2 vote_count Eşiği Kararı", 2)
add_para(doc,
    "Veri kalitesini sağlamak için oy sayısı (vote_count) filtresi uygulanmıştır. "
    "Bu parametre kalite ölçüsü değil, bilinirlik/popülerlik göstergesidir; ancak "
    "filtresiz çekim çöp doludur (örn. Documentary için 220.948 girdi: kısa film, "
    "TV yayını, postersiz, amatür içerik). Eşik değerinin nadir türlere etkisi:"
)
add_table(doc,
    ["Tür", "≥0", "≥5", "≥10 (seçilen)", "≥30"],
    [
        ["History",     "21.910",  "5.712",  "3.850", "2.149"],
        ["Documentary", "220.948", "17.450", "8.485", "2.621"],
        ["Animation",   "70.028",  "11.364", "7.663", "4.021"],
        ["Mystery",     "26.302",  "8.670",  "6.359", "3.572"],
        ["Fantasy",     "29.489",  "8.569",  "6.377", "3.858"],
    ],
    col_widths=[4, 3, 3, 4, 3]
)
add_para(doc,
    "Eşik 30 alındığında History yalnızca 2.149 filmde kalıyor ve 3.000 hedefine "
    "ulaşılamıyor; eşik 10'a düşürüldüğünde tüm nadir türler ≥3.850 adayla 3.000 "
    "hedefine rahatça ulaşabilmektedir. Bu nedenle vote_count ≥ 10 seçilmiştir."
)

# ── 4. Dengeli Örnekleme ────────────────────────────────────────────────────

add_heading(doc, "4. Dengeli Veri Örnekleme Algoritması", 1)

add_heading(doc, "4.1 v1 Veri Kümesindeki Dengesizlik", 2)
add_para(doc,
    "v1 veri kümesi (HTML kazıma ile toplanan 16.307 film) ileri derecede dengesizdi; "
    "en sık tür Drama (6.381 film), en nadir tür History (858 film) ile arasındaki "
    "oran yaklaşık 7,4 kat idi. Bu dengesizlik v1 modellerinde ciddi sorunlara yol açmıştır:"
)
add_table(doc,
    ["Tür", "Film Sayısı"],
    [
        ["Drama (maks)", "6.381"],
        ["Comedy", "5.116"],
        ["Thriller", "2.948"],
        ["Action", "2.615"],
        ["Adventure / Documentary", "2.017"],
        ["Horror", "2.004"],
        ["Crime", "2.128"],
        ["Romance", "2.269"],
        ["Science Fiction", "1.466"],
        ["Family", "1.698"],
        ["Fantasy", "1.334"],
        ["Animation", "1.266"],
        ["Mystery", "1.168"],
        ["History (min)", "858"],
    ],
    col_widths=[8, 4]
)
add_para(doc,
    "Dengesizliğin sonuçları: v1 model film başına 5–7 tür tahmin ediyordu (gerçek ort. ~2.3), "
    "en iyi türler Drama (F1=0.63) ve Comedy (F1=0.62) — frekans biası; "
    "en kötü türler History (F1=0.10) ve Documentary (F1=0.11) — az veri. "
    "Kök neden: veri dengesizliği + bunu telafi etmek için kullanılan ağır pos_weight (~5x) "
    "modeli aşırı tahmine (false positive'e) zorladı."
)

add_heading(doc, "4.2 İki Aşamalı Dengeli Örnekleme Yöntemi", 2)
add_para(doc,
    "v2 için tür bazında dengeli ve kombinasyon-farkında bir veri kümesi oluşturmak "
    "amacıyla iki aşamalı bir algoritma geliştirilmiştir:"
)
add_para(doc, "Aşama 1 — Aday Havuzu Oluşturma:", bold=True, space_after=2)
add_para(doc,
    "Her hedef tür için /discover/movie?with_genres=<id>&vote_count.gte=10 sorgusuyla "
    "sayfalanmış çekim yapılmış ve poster_path alanı olan filmler aday havuzuna alınmıştır. "
    "Havuz boyutu hedefin 1,5 katı (pool_factor=1.5) olarak belirlenmiş; toplamda ~37.000 "
    "benzersiz aday film elde edilmiştir."
)
add_para(doc, "Aşama 2 — Eksiklik Güdümlü Açgözlü Seçim:", bold=True, space_after=2)
add_para(doc,
    "Her adımda türler mevcut sayılarına göre sıralanır ve en eksik tür belirlenir. "
    "Bu türü içeren ve en az sayıda tür barındıran film (kombinasyon sadeliği) seçilir. "
    "Drama ve Comedy gibi baskın türler 'yolcu' olarak en sona bırakılarak overshoot önlenir. "
    "Sıralama anahtarı: (tür_sayısı, baskın_tür_cezası, film_id)."
)
add_para(doc,
    "Bu yöntem ayrı bir with_genres=a,b (AND) sorgusu gerektirmeden nadir türleri hedefe "
    "taşırken kombinasyon çeşitliliğini doğal olarak kapsıyor; co-occurrence dengesini "
    "de gözetiyor."
)

add_heading(doc, "4.3 Nihai Sonuçlar", 2)
try_add_figure(
    doc,
    "dist_before_after.png",
    "Şekil 1. Tür dağılımı: çekim öncesi (v1, n≈16.307) ve sonrası (v2, n=23.640)."
)
add_table(doc,
    ["Tür", "v2 Film Sayısı", "v1 Film Sayısı", "Değişim"],
    [
        ["History",             "3.000", "858",   "+2.142 (%250)"],
        ["Mystery",             "3.000", "1.168", "+1.832 (%157)"],
        ["Animation",           "3.000", "1.266", "+1.734 (%137)"],
        ["Fantasy",             "3.001", "1.334", "+1.667 (%125)"],
        ["Documentary",         "3.000", "2.017", "+983 (%49)"],
        ["Horror / Crime / vb.","3.000", "~2.000","~+1.000"],
        ["Action",              "3.762", "2.615", "+1.147"],
        ["Thriller",            "3.852", "2.948", "+904"],
        ["Comedy",              "4.178", "5.116", "-938 (yolcu kısıtı)"],
        ["Drama",               "6.658", "6.381", "+277 (yolcu kısıtı)"],
    ],
    col_widths=[5, 4, 4, 4]
)
add_para(doc,
    "Nihai veri kümesi 23.640 film içerir. Film başına ortalama 2,18 tür düşmektedir "
    "ve poster kapsamı %100'dür. Dengesizlik oranı 7,4 kattan 2,22 kata düşürülmüştür "
    "(nadir/orta türlerin tamamı ~3.000 örneğe çıkarılmıştır). Çekim süresi: HTML "
    "kazımada 3 saatte ~500 film + ban iken, API ile tüm veri seti tek oturumda ve "
    "ban almadan toplanmıştır."
)

# ── 5. Ön İşleme + CV ───────────────────────────────────────────────────────

add_heading(doc, "5. Veri Temizleme ve Çapraz Doğrulama Hazırlığı", 1)

add_heading(doc, "5.1 Temizleme Adımları", 2)
add_para(doc,
    "labels_v2.csv dosyasındaki 23.640 film aşağıdaki adımlarla temizlenmiştir:"
)
steps = [
    "Tür listesi boş olan satırlar kaldırıldı (yok).",
    "15 hedef tür dışında kalan türler (Western, War, TV Movie, Music) etiket vektöründen silindi.",
    "Poster dosyası bulunamayan veya bozuk olan satırlar düşürüldü.",
    "Temizleme sonucu: 0 satır düşürüldü — veri kümesi 23.640 filmde kaldı (%100 sağlam).",
    "MultiLabelBinarizer (15 sınıf) uygulanarak ikili etiket matrisi (N×15) oluşturuldu; mlb.pkl olarak kaydedildi.",
]
for s in steps:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(s).font.size = Pt(11)
doc.add_paragraph()

add_heading(doc, "5.2 5-Fold Çapraz Doğrulama Bölümlemesi", 2)
try_add_figure(
    doc,
    "cooccurrence_v2.png",
    "Şekil 2. v2 veri kümesinde türler arası birlikte-geçiş (co-occurrence) matrisi (n=23.640)."
)
add_para(doc,
    "Veri kümesi MultilabelStratifiedKFold (n_splits=5, seed=42) ile 5 kata bölünmüştür. "
    "Çoklu-etiket dağılımını koruyan bu yöntem, her fold'da tüm türlerin orantılı "
    "temsil edilmesini sağlar."
)
add_table(doc,
    ["Özellik", "Değer"],
    [
        ["Fold sayısı", "5"],
        ["Eğitim (fold başına)", "~18.900 film"],
        ["Doğrulama (fold başına)", "~4.730 film"],
        ["Train ∩ Val örtüşme", "0 (sıfır)"],
        ["Her film kaç kez val'da?", "Tam 1 kez"],
        ["Stratifikasyon", "Mükemmel (Drama %28.1–28.3, Action %15.9–16.0, History %12.7–12.8)"],
        ["Değerlendirme yöntemi", "OOF (out-of-fold): 5 fold val tahminleri birleştirilerek"],
    ],
    col_widths=[7, 9]
)
add_para(doc,
    "Çapraz doğrulama bütünlük kontrolü: fold'ların birleşimindeki benzersiz film sayısı "
    "= 23.640 (veri kümesi boyutu ile aynı) — örtüşme sıfır, tam kapsam doğrulandı."
)

# ── 6. Co-occurrence Riski ──────────────────────────────────────────────────

add_heading(doc, "6. Kombinasyon Dengesi ve Co-occurrence Riski", 1)
add_para(doc,
    "Çoklu-etiket veri kümelerinde sık birlikte görülen türler (örn. Comedy–Romance) "
    "modelin görsel öğrenme yerine etiket korelasyonunu sömürme riskini doğurur. "
    "Bu 'co-occurrence cheat' riski şu önlemlerle azaltılmıştır:"
)
measures = [
    "Veri toplama algoritması kombinasyon (ikili tür) dengesine dikkat ederek "
    "baskın eşleşmelerin aşırı birikmesini engellemiştir.",
    "Sonuç veri kümesinde nadir türlerin (History, Mystery vb.) F1 değerleri "
    "final modelde dramatik biçimde yükselmiştir; bu, saf korelasyon sömürüsünden "
    "beklenmeyen bir artıştır.",
    "Random baseline (film başına ~2.4 rasgele tür): macro-F1 = 0.149; "
    "frekans-öncül baseline (görsel yok): 0.146. En iyi transformer modeli (Swin) "
    "0.562 ile şansın ~3.8 katını başarmıştır.",
]
for m in measures:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(m).font.size = Pt(11)
doc.add_paragraph()
add_para(doc,
    "Bu risk tamamen ortadan kaldırılmış değildir; raporda açık şekilde not düşülmüştür."
)

# ── 7. Özet Tablo ──────────────────────────────────────────────────────────

add_heading(doc, "7. Özet: v1 → v2 Karşılaştırması", 1)
add_table(doc,
    ["Boyut", "v1 (HTML kazıma)", "v2 (TMDB API)"],
    [
        ["Film sayısı", "~16.307", "23.640"],
        ["Poster kapsamı", "~%91", "%100"],
        ["Dengesizlik (maks/min)", "~7.4x", "2.22x"],
        ["En nadir tür (History)", "858", "3.000"],
        ["Çekim hızı", "~500 film/3 saat + ban", "Tümü tek oturumda"],
        ["HTTP 429 / ban", "9.619 adet, 864 kayıp", "0 (ban yok)"],
        ["vote_count filtresi", "Yok", "≥10"],
        ["Kombinasyon dengesi", "Hayır", "Evet (greedy algoritma)"],
        ["Fold bölümlemesi", "70/15/15 random split", "5-fold stratified CV"],
    ],
    col_widths=[6, 5, 5]
)

# ── 8. Sonuç ───────────────────────────────────────────────────────────────

add_heading(doc, "8. Sonuç", 1)
add_para(doc,
    "Bu raporda film afişi veri kümesinin oluşturulma süreci iki ana aşamada "
    "incelenmiştir. HTML kazıma yaklaşımı ban sorunları ve kod hataları nedeniyle "
    "yeterli ve dengeli veri üretememektedir. TMDB resmî API'sine geçişle birlikte "
    "23.640 filmlik, %100 poster kapsamlı ve tür dengesizliği 7.4x'ten 2.22x'e "
    "düşürülmüş bir veri kümesi elde edilmiştir."
)
add_para(doc,
    "İki aşamalı dengeli örnekleme algoritması nadir türleri (History, Mystery vb.) "
    "hedef sayıya taşırken baskın türlerin aşırı şişmesini kontrol altında tutmuştur. "
    "MultilabelStratifiedKFold ile gerçekleştirilen 5-fold bölümlemesi her türün "
    "orantılı temsil edildiğini güvence altına almıştır."
)
add_para(doc,
    "Bu veri hazırlığı sayesinde modelin frekans biasından değil görsel sinyalden "
    "öğrenmesi sağlanmıştır: final modelde (Swin) nadir türlerin F1 değerleri "
    "dramatik biçimde yükselmiş ve random baseline'ın ~3.8 katı performans elde edilmiştir."
)

# ── kaydet ──────────────────────────────────────────────────────────────────

doc.save(OUT)
print(f"Kaydedildi: {OUT}")
